"""Regression tests for the three reader bugs found in the v2 audit
(docs/refinement/01_AUDIT_REPORT.md, section 1)."""
from __future__ import annotations

import os

import openpyxl
import pytest
import seed_history
import transform
from common import COUNTRIES, TRACERS
from load import connect, init_db, load_return

# The 20 first-round country names exactly as PEN_PLUS_MONITORING.xlsx spells
# them on its "PEN plus phases" sheet, including the two Congo variants that
# matched nothing before the aliases were added.
WORKBOOK_NAMES = [
    "Benin", "Burkina Faso", "Cameroon", "Democratic Republic of Congo",
    "Congo, Republic of", "Ethiopia", "Ghana", "Kenya", "Liberia", "Lesotho",
    "Mozambique", "Malawi", "Niger", "Nigeria", "Rwanda", "Sierra Leone",
    "Tanzania", "Uganda", "Zambia", "Zimbabwe",
]


def _real_layout_workbook(path):
    """The real sheet's shape: banner rows 1-6, 'Countries' in column C of
    row 7, a row number in column B, step headers split over rows 7-8."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PEN plus phases"
    ws.cell(row=7, column=3, value="Countries")
    ws.cell(row=7, column=4, value="FHASE 1. Assessment")
    ws.cell(row=8, column=4, value="Step 1: Conduct a comprehensive assessment")
    ws.cell(row=8, column=5, value="Step 2. SWOT Analysis")
    for i, name in enumerate(WORKBOOK_NAMES):
        ws.cell(row=9 + i, column=2, value=str(i + 1))
        ws.cell(row=9 + i, column=3, value=name)
        ws.cell(row=9 + i, column=4, value="Yes")
        ws.cell(row=9 + i, column=5, value="en cours")
    wb.save(path)


def test_phases_reader_finds_all_20_countries_in_column_c(tmp_path):
    path = str(tmp_path / "monitoring.xlsx")
    _real_layout_workbook(path)
    phases = seed_history.from_monitoring_phases(path)
    assert len(phases) == 20
    assert {"COD", "COG", "TZA", "KEN"} <= set(phases)
    assert phases["COD"][1] == "yes"
    assert phases["COG"][2] == "under_development"


def test_monitoring_aliases_resolve():
    assert COUNTRIES["Democratic Republic of Congo"][0] == "COD"
    assert COUNTRIES["Congo, Republic of"][0] == "COG"
    assert COUNTRIES["KEN Kenya"][0] == "KEN"


REAL_MONITORING = os.path.join(os.path.dirname(__file__), "..", "..", "data", "raw",
                               "PEN_PLUS_MONITORING.xlsx")


@pytest.mark.skipif(not os.path.exists(REAL_MONITORING),
                    reason="raw workbook not in the repository (data/raw is git-ignored)")
def test_phases_reader_reads_20_countries_from_the_real_workbook():
    assert len(seed_history.from_monitoring_phases(REAL_MONITORING)) == 20


def test_severe_htn_is_not_a_sentinel_condition():
    assert "severe_htn" not in TRACERS
    assert TRACERS == ["t1d", "scd", "rhd"]


def test_severe_htn_never_enters_the_sentinel_total(tmp_path, make_rec):
    """Uganda's slide: SCD+T1D+RHD = 2,226 plus severe HTN 15. The total must
    be 2,226; the 15 stays visible as its own condition row."""
    db_path = str(tmp_path / "store.db")
    con = init_db(db_path)
    rec = make_rec(country_name="Uganda", period_id="2026-Q1", quarter_id="2026-Q1")
    rec["patient_stock"] = [
        {"condition": "scd", "ever_enrolled": 2004, "active_end": None},
        {"condition": "t1d", "ever_enrolled": 206, "active_end": None},
        {"condition": "rhd", "ever_enrolled": 16, "active_end": None},
        {"condition": "severe_htn", "ever_enrolled": 15, "active_end": None},
    ]
    load_return(con, rec, verdict="accepted")
    con.close()
    transform.build(db_path)
    gold = {(g["indicator_code"], g["disagg_key"], g["disagg_value"]): g["numerator"]
            for g in connect(db_path).execute("SELECT * FROM gold_indicator WHERE iso3='UGA'")}
    assert gold[("2.6", "all", "all")] == 2226
    assert gold[("2.5", "condition", "severe_htn")] == 15


def _icppa_workbook(path, site_names):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enrolment_by_Year"
    ws.append(["Country", "Reporting unit / partner", "Year", "Include in country total",
               "Sickle cell disease", "Value type"])
    ws.append(["Zimbabwe", "National", 2025, "Y", 10, "flow"])
    st = wb.create_sheet("Staff_Trained_by_Year")
    st.append(["Country", "Year", "Include in country total"])
    si = wb.create_sheet("Sites_and_Clinics")
    si.append(["Country", "Number of PEN-Plus sites / clinics", "Site or district names",
               "Supporting partners and sites each", "Status / stage"])
    si.append(["Zimbabwe", 8, site_names, "WHO", "Scale-up"])
    wb.save(path)


def test_icppa_reader_does_not_build_facilities_from_site_text(tmp_path):
    """'Rushinga, Mutoko, Tsholotsho' used to become one facility, and
    'planned: Kambia and Moyamba' an operational one."""
    path = str(tmp_path / "icppa.xlsx")
    _icppa_workbook(path, "Rushinga, Mutoko, Tsholotsho; planned: Kambia and Moyamba")
    recs = seed_history.from_icppa(path)
    assert recs, "the enrolment row should still produce a return"
    for rec in recs.values():
        assert rec["facilities"] == []


def test_aliases_never_become_the_display_name(tmp_path):
    """The first name listed for an ISO3 is what dim_country shows; an alias
    placed first once made the site read "KEN Kenya"."""
    con = init_db(str(tmp_path / "store.db"))
    names = {r["iso3"]: r["name"] for r in con.execute("SELECT iso3, name FROM dim_country")}
    assert len(names) == 31
    assert names["KEN"] == "Kenya"
    assert names["COD"] == "Democratic Republic of the Congo"
    assert names["COG"] == "Republic of the Congo"
