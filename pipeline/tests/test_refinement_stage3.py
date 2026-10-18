"""Stage 3 of refinement v2: the ICPPA v2 reading rules, the curated
reference files, and the validation checks that decide what is held."""
from __future__ import annotations

import baseline
import openpyxl
import seed_history
import transform
from fixtures.build_national_v4_return import DEFAULT, build_national_v4_return
from load import connect, init_db, load_return
from parse import parse_return
from validate import run_checks

ENROL_HEAD = ["Country", "Reporting unit / partner", "Year", "Include in country total",
              "Sickle cell disease", "Type 1 diabetes", "Rheumatic heart disease",
              "Severe hypertension", "Source document", "Slide no.", "Value type"]
DECK = "Slide_Deck_-_Day_3_-_Main_room.pptx"


def _workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enrolment_by_Year"
    ws.append(ENROL_HEAD)
    # flows: ever enrolled is the running sum; severe HTN never in the sentinel total
    ws.append(["Uganda", "UINCD", 2025, "Y", 100, 10, 1, 15, DECK, 254, "flow"])
    ws.append(["Uganda", "UINCD", "2026 (to end March)", "Y", 50, 5, 1, 0, DECK, 254, "flow"])
    # cumulative: a stock, never summed; earlier rows marked N still load as their own year
    ws.append(["Kenya", "National", 2025, "N", 213, 227, 19, None, DECK, 93, "cumulative"])
    ws.append(["Kenya", "National", 2026, "Y", 229, 310, 20, None, DECK, 93, "cumulative"])
    # active: Rwanda's series is active patients, flagged N, loaded as active
    ws.append(["Rwanda", "National", 2024, "N", None, 2847, None, None, DECK, 51, "active"])
    # unclear: never loaded
    ws.append(["Malawi", "PIH", 2025, "N", 877, 620, 146, None, DECK, 80, "unclear"])
    st = wb.create_sheet("Staff_Trained_by_Year")
    st.append(["Country", "Year", "Include in country total", "Medical doctors",
               "Nurses (incl. physician assistants)", "Source document", "Slide no.", "Value type"])
    st.append(["Uganda", 2025, "Y", 1, 2, DECK, 255, "flow"])
    st.append(["Uganda", 2026, "Y", 1, 0, DECK, 255, "flow"])
    si = wb.create_sheet("Sites_and_Clinics")
    si.append(["Country", "Number of PEN-Plus sites / clinics"])
    lm = wb.create_sheet("Load_Map_Other_Figures")
    lm.append(["Country", "Load as", "Indicator", "Condition", "Value", "Value type", "As of",
               "Load?", "Source document", "Slide no.", "Why"])
    lm.append(["Zambia", "2.5 all PEN-Plus conditions", "2.5", "all", 1397, "cumulative",
               "2026-06-30", "Y", DECK, "200, 201", ""])
    lm.append(["Zambia", "2.6 ever enrolled", "2.6", "cardiac, not split", 220, "cumulative",
               "2026-06-30", "Y", DECK, "200, 201", ""])
    lm.append(["Zambia", "2.6 active in care", "2.6", "all", 1328, "active", "2026-06-30", "Y",
               DECK, "200, 201", ""])
    lm.append(["Ghana", "3.3 health workers trained", "3.3", None, 426, "cumulative",
               "2026-06-30", "N", DECK, 57, "not loaded"])
    wb.save(path)


def _build(tmp_path):
    path = str(tmp_path / "icppa.xlsx")
    _workbook(path)
    skipped = []
    recs = seed_history.from_icppa(path, skipped)
    db_path = str(tmp_path / "store.db")
    con = init_db(db_path)
    for rec in recs.values():
        load_return(con, rec, source_kind="historical", provenance="test")
    con.close()
    transform.build(db_path)
    gold = {(g["iso3"], g["period_id"], g["indicator_code"], g["disagg_key"], g["disagg_value"]): g
            for g in connect(db_path).execute("SELECT * FROM gold_indicator")}
    return gold, skipped


def test_flow_rows_accumulate_and_severe_htn_stays_out(tmp_path):
    gold, _ = _build(tmp_path)
    assert gold[("UGA", "2025-A", "2.6", "all", "all")]["numerator"] == 111
    assert gold[("UGA", "2026-A", "2.6", "all", "all")]["numerator"] == 111 + 56
    assert gold[("UGA", "2026-A", "2.5", "condition", "severe_htn")]["numerator"] == 15
    assert gold[("UGA", "2026-A", "2.6", "all", "all")]["source_ref"] == "D3M p254"


def test_cumulative_rows_are_stocks_never_summed(tmp_path):
    gold, _ = _build(tmp_path)
    assert gold[("KEN", "2025-A", "2.6", "all", "all")]["numerator"] == 459
    assert gold[("KEN", "2026-A", "2.6", "all", "all")]["numerator"] == 559
    assert gold[("KEN", "2026-A", "2.6", "condition", "t1d")]["numerator"] == 310


def test_active_series_loads_as_active_and_unclear_is_skipped(tmp_path):
    gold, skipped = _build(tmp_path)
    assert gold[("RWA", "2024-A", "2.6a", "all", "all")]["numerator"] == 2847
    assert not any(k[0] == "MWI" for k in gold)
    assert any(c == "Malawi" and "unclear" in reason for c, _, reason in skipped)


def test_load_map_rows(tmp_path):
    gold, skipped = _build(tmp_path)
    assert gold[("ZMB", "2026-A", "2.5", "all", "all")]["numerator"] == 1397
    # Zambia's active figure covers all conditions: published with 2.5, not as sentinel 2.6a
    assert gold[("ZMB", "2026-A", "2.5", "measure", "active")]["numerator"] == 1328
    assert ("ZMB", "2026-A", "2.6a", "all", "all") not in gold or \
        gold[("ZMB", "2026-A", "2.6a", "all", "all")]["numerator"] is None
    # "cardiac, not split" counts in the sentinel total as the country reported it
    assert gold[("ZMB", "2026-A", "2.6", "all", "all")]["numerator"] == 220
    assert ("GHA", "2026-A", "3.3", "all", "all") not in gold   # Load? = N


def test_trained_to_date_is_the_running_sum(tmp_path):
    gold, _ = _build(tmp_path)
    assert gold[("UGA", "2025-A", "3.3", "all", "all")]["numerator"] == 3
    assert gold[("UGA", "2026-A", "3.3", "all", "all")]["numerator"] == 1
    assert gold[("UGA", "2026-A", "3.3", "measure", "to_date")]["numerator"] == 4


def test_query_seed_loads_all_28_and_closes_only_what_this_build_fixed(tmp_path):
    con = init_db(str(tmp_path / "store.db"))
    seeded, closed = baseline.load_queries(con)
    assert seeded == 28
    assert closed == len(baseline.RESOLVED)
    status = {q["query_id"]: q["status"] for q in con.execute("SELECT * FROM fact_query")}
    assert status["Q-ZWE-01"] == "open"         # held until MoHCC confirms
    # programme decision (18 Oct 2026): the regional figure is published as the
    # sum of country-reported figures with its coverage, so Q-REG-01 is closed
    assert status["Q-REG-01"] == "closed"
    assert status["Q-REG-02"] == "closed"       # severe HTN removed by this build
    high_open = [q for q in con.execute(
        "SELECT * FROM fact_query WHERE severity='High' AND status='open'")]
    assert {q["query_id"] for q in high_open} == {"Q-ZWE-01", "Q-MWI-01"}


def test_facility_register_counts_only_active_facilities_for_2_3(tmp_path):
    con = init_db(str(tmp_path / "store.db"))
    assert baseline.load_facility_register(con) == 43
    counts = dict(con.execute(
        "SELECT register_status, COUNT(*) FROM dim_facility GROUP BY register_status").fetchall())
    assert counts == {"active": 33, "planned": 6, "to_confirm": 3, "partial": 1}
    operational = con.execute(
        "SELECT COUNT(*) FROM dim_facility WHERE status='operational'").fetchone()[0]
    assert operational == 33


def test_a_return_with_active_above_ever_fails_r04_and_is_held(tmp_path):
    """DATA_QUALITY_SPEC.md section 5: an arithmetic error appears as a failed
    R04 and as a held figure, outside every total."""
    path = str(tmp_path / "GHA_2026_Q3.docx")
    bad = dict(DEFAULT["sentinel_26"])
    bad["Type 1 diabetes"] = (300, 20, 450, 0)          # active 450 > ever 300
    build_national_v4_return(path, sentinel_26=bad)
    db_path = str(tmp_path / "store.db")
    con = init_db(db_path)
    load_return(con, parse_return(path), verdict="accepted")
    failures = run_checks(con)
    con.close()
    assert failures["R04"] == 1
    transform.build(db_path)
    check = connect(db_path)
    r04 = check.execute("SELECT * FROM fact_check WHERE rule='R04'").fetchone()
    assert r04["outcome"] == "failed" and "active 450 > ever 300" in r04["detail"]
    g = check.execute("SELECT * FROM gold_indicator WHERE indicator_code='2.6'"
                      " AND disagg_key='all'").fetchone()
    assert g["held"] == 1 and "V-R04-GHA-2026-Q3" in g["held_by"]
    q = check.execute("SELECT * FROM fact_query WHERE rule='R04'").fetchone()
    assert q["severity"] == "High" and q["status"] == "open"


def test_a_clean_new_tool_return_passes_every_testable_rule(tmp_path):
    path = str(tmp_path / "GHA_2026_Q3.docx")
    build_national_v4_return(path)
    con = init_db(str(tmp_path / "store.db"))
    load_return(con, parse_return(path), verdict="accepted")
    failures = run_checks(con)
    assert not any(failures.values()), failures
    outcomes = dict(con.execute("SELECT rule, outcome FROM fact_check").fetchall())
    for rule in ("R01", "R03", "R04", "R05", "R07", "R08", "R09", "R11", "R12", "R13", "R16"):
        assert outcomes[rule] == "passed", rule
