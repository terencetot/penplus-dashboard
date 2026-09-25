"""consolidate.py: the partner workbook and per-country data-quality report."""
from __future__ import annotations

import os

from consolidate import all_country_reports, build_workbook, country_report
from load import load_return


def test_country_report_no_return_on_file(db):
    report = country_report(db, "GHA")
    assert "No return on file" in report


def test_country_report_first_return_has_no_comparison(db, make_rec):
    load_return(db, make_rec(), verdict="accepted")
    report = country_report(db, "GHA")
    assert "Ghana (GHA)" in report
    assert "No prior accepted return to compare against" in report
    assert "No action needed" in report


def test_country_report_flags_self_reported_divergence(db, make_rec):
    rec = make_rec()
    rec["quality"]["reported_completeness_pct"] = 40  # computed is 80% (8 of 10)
    load_return(db, rec, verdict="accepted")
    report = country_report(db, "GHA")
    assert "percentage points apart" in report


def test_country_report_flags_reconciliation_no(db, make_rec):
    rec = make_rec()
    rec["reconciliation"] = {"annex_a_vs_2_3": "No"}
    load_return(db, rec, verdict="accepted")
    report = country_report(db, "GHA")
    assert "DOES NOT RECONCILE" in report


def test_country_report_high_severity_query_recommends_contact(db, make_rec):
    rid = load_return(db, make_rec(), verdict="accepted")
    db.execute(
        "INSERT INTO query_register(return_id,severity,section,field,observed,expected,"
        "question,status,raised_at) VALUES (?,?,?,?,?,?,?,'open',?)",
        (rid, "High", "2.6", "Retention denominator", "obs", "exp", "Please check.", "2026-04-01"))
    report = country_report(db, "GHA")
    assert "Contact the country" in report


def test_country_report_growth_more_than_doubled_is_flagged(db, make_rec):
    load_return(db, make_rec(), verdict="accepted")  # t1d ever_enrolled=120
    rec2 = make_rec(period_id="2026-Q2", quarter_id="2026-Q2")
    rec2["patient_stock"][0]["ever_enrolled"] = 300
    load_return(db, rec2, verdict="accepted")
    report = country_report(db, "GHA")
    assert "more than doubled" in report


def test_all_country_reports_writes_one_file_per_country(db, make_rec, tmp_path):
    load_return(db, make_rec(), verdict="accepted")
    load_return(db, make_rec(country_name="Kenya", period_id="2026-Q1", quarter_id="2026-Q1"),
                verdict="accepted")
    out_dir = str(tmp_path / "reports")
    written = all_country_reports(db, out_dir)
    names = {os.path.basename(p) for p in written}
    assert "GHA.md" in names
    assert "KEN.md" in names


def test_build_workbook_creates_three_sheets(db, make_rec, tmp_path):
    load_return(db, make_rec(), verdict="accepted")
    out_path = str(tmp_path / "partner_workbook.xlsx")
    build_workbook(db, out_path)
    assert os.path.exists(out_path)

    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    assert wb.sheetnames == ["Indicators by country", "Data quality", "Open queries"]
    quality_sheet = wb["Data quality"]
    header = [c.value for c in quality_sheet[1]]
    assert "Country" in header and "Completeness %" in header
